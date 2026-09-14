from io import BytesIO
from typing import Any, Iterator, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

_ACCENT = "1F8F63"
_ACCENT_DARK = "16674A"
_TEXT = "1B241F"
_MUTED = "66766B"
_STRIPE = "F3F6F2"
_LINE = "D8E0DA"
_SUCCESS_BG = "D6F2E3"
_SUCCESS_TEXT = "14603F"
_FAILURE_BG = "FBE0E0"
_FAILURE_TEXT = "9B1C1C"

_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_HEADING_FONT = Font(name="Calibri", size=16, bold=True, color=_ACCENT_DARK)
_TEXT_FONT = Font(name="Calibri", size=11, color=_TEXT)
_MUTED_FONT = Font(name="Calibri", size=11, color=_MUTED)
_BOLD_FONT = Font(name="Calibri", size=11, bold=True, color=_TEXT)
_SUCCESS_FONT = Font(name="Calibri", size=11, bold=True, color=_SUCCESS_TEXT)
_FAILURE_FONT = Font(name="Calibri", size=11, bold=True, color=_FAILURE_TEXT)

_HEADER_FILL = PatternFill("solid", fgColor=_ACCENT)
_STRIPE_FILL = PatternFill("solid", fgColor=_STRIPE)
_SUCCESS_FILL = PatternFill("solid", fgColor=_SUCCESS_BG)
_FAILURE_FILL = PatternFill("solid", fgColor=_FAILURE_BG)

_SIDE = Side(style="thin", color=_LINE)
_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)

_TEXT_ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

_MIN_COLUMN_WIDTH = 16
_MAX_COLUMN_WIDTH = 60
_TEXT_COLUMN_WIDTH = 100
_LABEL_COLUMN_WIDTH = 34
_VALUE_COLUMN_WIDTH = 16
_HEADER_ROW_HEIGHT = 30
_HEADING_ROW_HEIGHT = 34
_ROW_HEIGHT = 22


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _set_widths(sheet: Worksheet, widths: Sequence[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _measure_widths(rows: Sequence[Sequence[str]], columns: int) -> list[int]:
    widths: list[int] = []
    for index in range(columns):
        longest = max((len(row[index]) for row in rows if index < len(row)), default=0)
        widths.append(min(max(longest + 6, _MIN_COLUMN_WIDTH), _MAX_COLUMN_WIDTH))
    return widths


def _create_sheet(workbook: Workbook, title: str) -> Worksheet:
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    return sheet


def _write_heading(sheet: Worksheet, heading: str) -> None:
    cell = sheet.cell(row=1, column=1, value=heading)
    cell.font = _HEADING_FONT
    cell.alignment = _TEXT_ALIGNMENT
    sheet.row_dimensions[1].height = _HEADING_ROW_HEIGHT


def read_rows(content: bytes, columns: int) -> Iterator[list[str]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        for values in workbook.worksheets[0].iter_rows(values_only=True):
            cells = [_cell_text(value) for value in values[:columns]]
            yield cells + [""] * (columns - len(cells))
    finally:
        workbook.close()


def create_workbook() -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    return workbook


def add_table_sheet(
    workbook: Workbook,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]] = (),
    widths: Sequence[float] | None = None,
) -> Worksheet:
    sheet = _create_sheet(workbook, title)

    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER_ALIGNMENT
        cell.border = _BORDER
    sheet.row_dimensions[1].height = _HEADER_ROW_HEIGHT

    for values in rows:
        sheet.append(list(values))
        number = sheet.max_row
        sheet.row_dimensions[number].height = _ROW_HEIGHT
        for cell in sheet[number]:
            cell.font = _TEXT_FONT
            cell.alignment = _TEXT_ALIGNMENT
            cell.border = _BORDER
            if number % 2 == 1:
                cell.fill = _STRIPE_FILL

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    _set_widths(
        sheet,
        widths
        if widths is not None
        else _measure_widths([list(headers), *rows], len(headers)),
    )

    return sheet


def add_fields_sheet(
    workbook: Workbook,
    title: str,
    heading: str,
    fields: Sequence[tuple[str, Any]],
) -> Worksheet:
    sheet = _create_sheet(workbook, title)
    _set_widths(sheet, (_LABEL_COLUMN_WIDTH, _VALUE_COLUMN_WIDTH))
    _write_heading(sheet, heading)

    for number, (label, value) in enumerate(fields, start=3):
        label_cell = sheet.cell(row=number, column=1, value=label)
        label_cell.font = _MUTED_FONT
        label_cell.alignment = _TEXT_ALIGNMENT

        value_cell = sheet.cell(row=number, column=2, value=value)
        value_cell.font = _BOLD_FONT
        value_cell.alignment = _CENTER_ALIGNMENT

        sheet.row_dimensions[number].height = _ROW_HEIGHT

    return sheet


def add_text_sheet(
    workbook: Workbook,
    title: str,
    heading: str,
    lines: Sequence[str],
) -> Worksheet:
    sheet = _create_sheet(workbook, title)
    _set_widths(sheet, (_TEXT_COLUMN_WIDTH,))
    _write_heading(sheet, heading)

    for number, line in enumerate(lines, start=3):
        cell = sheet.cell(row=number, column=1, value=f"•   {line}")
        cell.font = _TEXT_FONT
        cell.alignment = _TEXT_ALIGNMENT
        sheet.row_dimensions[number].height = _ROW_HEIGHT

    return sheet


def add_column_choices(
    sheet: Worksheet,
    column: int,
    choices: Sequence[str],
    rows: int,
) -> None:
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(choices) + '"',
        allow_blank=True,
    )
    sheet.add_data_validation(validation)

    letter = get_column_letter(column)
    validation.add(f"{letter}2:{letter}{rows + 1}")


def apply_success_style(cell: Cell) -> None:
    cell.font = _SUCCESS_FONT
    cell.fill = _SUCCESS_FILL
    cell.alignment = _CENTER_ALIGNMENT


def apply_failure_style(cell: Cell) -> None:
    cell.font = _FAILURE_FONT
    cell.fill = _FAILURE_FILL
    cell.alignment = _CENTER_ALIGNMENT


def to_bytes(workbook: Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
